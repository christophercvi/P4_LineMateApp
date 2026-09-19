"""Render LineMate seed and sample documents into real office files.

Input:  back-end/seed_data/documents.json (metadata, fileName, fileKind)
        back-end/seed_sources/seed/<DOC-ID>.json    (full Markdown body)
        back-end/seed_sources/samples/SMP-xx.json   (sample documents for end users to upload)
        back-end/seed_sources/photos/*              (source photos for ticket attachments)
Output: back-end/seed_documents/<fileName>          (loaded at startup)
        back-end/sample_documents/<slug>.<ext>      (for people to try the upload screens)
        back-end/seed_attachments/<name>            (ticket attachment files)

PDFs are produced by writing a DOCX and converting it with headless LibreOffice so they carry a
real text layer. The incident JPG is a rendered delivery slip so OCR (tesseract) has text to read.
Run from the back-end folder:  uv run python scripts/render_seed_files.py
"""

import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PPt

BACK = Path(__file__).resolve().parents[1]
EXPANDED = Path(sys.argv[1]) if len(sys.argv) > 1 else BACK / "seed_sources"
SEED_OUT = BACK / "seed_documents"
SAMPLE_OUT = BACK / "sample_documents"
ATT_OUT = BACK / "seed_attachments"
SOURCE_PHOTOS = BACK / "seed_sources" / "photos"
FONT = "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"
ACCENT = RGBColor(0xC2, 0x41, 0x0C)

# ------------------------------------------------------------------ markdown -> blocks


def parse(md: str) -> list[tuple]:
    blocks: list[tuple] = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)$", s):
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif s.startswith(">"):
            blocks.append(("quote", s.lstrip("> ").strip()))
            i += 1
        elif re.match(r"^\s*([-*]|\d+\.)\s+", line):
            items = []
            ordered = bool(re.match(r"^\d+\.", s))
            while i < len(lines) and (m := re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])):
                depth = 1 if len(m.group(1)) >= 2 else 0
                items.append((depth, m.group(2)[0].isdigit(), m.group(3).strip()))
                i += 1
            blocks.append(("list", ordered, items))
        else:
            para = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\||>|\s*([-*]|\d+\.)\s)", lines[i]):
                para.append(lines[i].strip())
                i += 1
            blocks.append(("p", " ".join(para)))
    return blocks


def plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`", lambda m: m.group(1) or m.group(2) or m.group(3), text)


def runs(text: str) -> list[tuple[str, bool, bool]]:
    out = []
    for part in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


def sections(blocks: list[tuple]) -> tuple[str, list[tuple[str, list[tuple]]]]:
    title, out, cur = "", [], None
    for b in blocks:
        if b[0] == "h" and b[1] == 1 and not title:
            title = plain(b[2])
            continue
        if b[0] == "h" and b[1] == 2:
            cur = (plain(b[2]), [])
            out.append(cur)
            continue
        if cur is None:
            cur = ("Overview", [])
            out.append(cur)
        cur[1].append(b)
    return title, out


# ------------------------------------------------------------------ writers


def write_docx(meta: dict, body: str, path: Path) -> None:
    doc = Docx()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    core = doc.core_properties
    core.title, core.author, core.subject = meta["title"], meta.get("ownerName", "LineMate"), meta.get("category", "")
    for b in parse(body):
        kind = b[0]
        if kind == "h":
            h = doc.add_heading(plain(b[2]), level=min(b[1], 3) if b[1] > 1 else 0)
            for r in h.runs:
                r.font.color.rgb = ACCENT if b[1] <= 2 else RGBColor(0x33, 0x33, 0x33)
        elif kind == "p":
            p = doc.add_paragraph()
            for t, bold, ital in runs(b[1]):
                r = p.add_run(t)
                r.bold, r.italic = bold, ital
        elif kind == "quote":
            p = doc.add_paragraph()
            r = p.add_run(plain(b[1]))
            r.italic = True
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif kind == "list":
            for depth, num, text in b[2]:
                style = ("List Number" if num else "List Bullet") + (" 2" if depth else "")
                p = doc.add_paragraph(style=style)
                for t, bold, ital in runs(text):
                    r = p.add_run(t)
                    r.bold, r.italic = bold, ital
        elif kind == "table":
            rows = b[1]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Light Grid Accent 2"
            for ri, row in enumerate(rows):
                for ci in range(width):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    for t, bold, ital in runs(row[ci] if ci < len(row) else ""):
                        r = para.add_run(t)
                        r.bold, r.italic = bold or ri == 0, ital
            doc.add_paragraph()
    foot = doc.sections[0].footer.paragraphs[0]
    foot.text = f"{meta['id']} · {meta['title']} · Hearthline Kitchen (illustrative content)"
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def write_pptx(meta: dict, body: str, path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title, secs = sections(parse(body))
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title or meta["title"]
    s.placeholders[1].text = meta.get("summary", "")
    for name, blocks in secs:
        lines: list[tuple[int, str]] = []
        for b in blocks:
            if b[0] in ("p", "quote"):
                lines.append((0, plain(b[1])))
            elif b[0] == "h":
                lines.append((0, plain(b[2])))
            elif b[0] == "list":
                lines += [(d, plain(t)) for d, _, t in b[2]]
            elif b[0] == "table":
                head, *rest = b[1] or [[]]
                for r in rest:
                    lines.append((0, " · ".join(f"{h}: {plain(c)}" for h, c in zip(head, r, strict=False) if c)))
        for part in range(0, max(1, len(lines)), 9):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = name + (" (cont.)" if part else "")
            tf = slide.placeholders[1].text_frame
            tf.clear()
            chunk = lines[part : part + 9] or [(0, "")]
            for idx, (depth, text) in enumerate(chunk):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text, p.level = text, depth
                for r in p.runs:
                    r.font.size = PPt(18 if depth == 0 else 16)
    prs.core_properties.title = meta["title"]
    prs.save(path)


def write_xlsx(meta: dict, body: str, path: Path) -> None:
    wb = Workbook()
    title, secs = sections(parse(body))
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = title or meta["title"]
    ws["A1"].font = Font(bold=True, size=14, color="C2410C")
    ws["A2"] = meta.get("summary", "")
    ws["A3"] = f"Station: {meta.get('station', '')} · Category: {meta.get('category', '')}"
    ws.column_dimensions["A"].width = 90
    used = {"Overview"}
    head_fill = PatternFill("solid", fgColor="FDE7D9")
    for name, blocks in secs:
        sheet = re.sub(r"[\[\]:*?/\\]", "", plain(name))[:28] or "Section"
        base, n = sheet, 2
        while sheet in used:
            sheet = f"{base[:25]} {n}"
            n += 1
        used.add(sheet)
        sh = wb.create_sheet(sheet)
        sh["A1"] = plain(name)
        sh["A1"].font = Font(bold=True, size=12, color="C2410C")
        row = 3
        widths: dict[int, int] = {}
        for b in blocks:
            if b[0] == "table":
                for ri, r in enumerate(b[1]):
                    for ci, c in enumerate(r, 1):
                        cell = sh.cell(row=row, column=ci, value=plain(c))
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        widths[ci] = max(widths.get(ci, 10), min(60, len(plain(c)) + 2))
                        if ri == 0:
                            cell.font, cell.fill = Font(bold=True), head_fill
                    row += 1
                row += 1
            elif b[0] == "list":
                for d, num, t in b[2]:
                    sh.cell(row=row, column=1 + d, value=("• " if not num else "") + plain(t)).alignment = Alignment(wrap_text=True)
                    widths[1] = max(widths.get(1, 10), 80)
                    row += 1
                row += 1
            elif b[0] in ("p", "quote", "h"):
                sh.cell(row=row, column=1, value=plain(b[1] if b[0] != "h" else b[2]))
                widths[1] = max(widths.get(1, 10), 80)
                row += 1
        for ci, w in widths.items():
            sh.column_dimensions[chr(64 + ci)].width = w
    wb.properties.title = meta["title"]
    wb.save(path)


def write_slip_jpg(meta: dict, body: str, path: Path) -> None:
    """A photographed-looking delivery slip / incident note with real text for OCR."""
    lines: list[tuple[str, bool]] = []
    for b in parse(body):
        if b[0] == "h":
            lines.append((plain(b[2]), True))
        elif b[0] in ("p", "quote"):
            lines.append((plain(b[1]), False))
        elif b[0] == "list":
            lines += [(("   - " if d else "- ") + plain(t), False) for d, _, t in b[2]]
        elif b[0] == "table":
            lines += [("  |  ".join(plain(c) for c in r), i == 0) for i, r in enumerate(b[1])]
    wrapped: list[tuple[str, bool]] = []
    for text, bold in lines[:70]:
        while len(text) > 78:
            cut = text.rfind(" ", 0, 78)
            cut = cut if cut > 20 else 78
            wrapped.append((text[:cut], bold))
            text = "   " + text[cut:].strip()
        wrapped.append((text, bold))
    wrapped = wrapped[:58]
    W, lh = 1700, 34
    H = 180 + lh * len(wrapped) + 120
    img = Image.new("RGB", (W, H), (250, 248, 242))
    d = ImageDraw.Draw(img)
    f, fb, ft = ImageFont.truetype(FONT, 25), ImageFont.truetype(FONT_BOLD, 25), ImageFont.truetype(FONT_BOLD, 40)
    d.rectangle([40, 40, W - 40, H - 40], outline=(60, 60, 60), width=3)
    d.text((80, 70), meta["title"], font=ft, fill=(30, 30, 30))
    y = 150
    for text, bold in wrapped:
        d.text((80, y), text, font=fb if bold else f, fill=(25, 25, 25))
        y += lh
    img = img.rotate(-0.6, expand=True, fillcolor=(235, 232, 225))
    img.save(path, "JPEG", quality=88)


def docx_to_pdf(docx_paths: list[Path], outdir: Path) -> None:
    if not docx_paths:
        return
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(outdir),
                *map(str, docx_paths),
            ],
            check=True,
            capture_output=True,
            timeout=600,
        )


def render(meta: dict, body: str, out: Path, pdf_jobs: list[tuple[Path, Path]], tmp: Path) -> None:
    kind = meta["fileKind"]
    out.parent.mkdir(parents=True, exist_ok=True)
    match kind:
        case "md":
            out.write_text(body.strip() + "\n", encoding="utf-8")
        case "docx":
            write_docx(meta, body, out)
        case "pptx":
            write_pptx(meta, body, out)
        case "xlsx":
            write_xlsx(meta, body, out)
        case "jpg":
            write_slip_jpg(meta, body, out)
        case "pdf":
            src = tmp / (out.stem + ".docx")
            write_docx(meta, body, src)
            pdf_jobs.append((src, out))
        case _:
            raise ValueError(kind)


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60]


def main() -> None:
    docs = json.loads((BACK / "seed_data" / "documents.json").read_text())
    crew = {c["id"]: c for c in json.loads((BACK / "seed_data" / "crew.json").read_text())}
    for d in (SEED_OUT, SAMPLE_OUT, ATT_OUT):
        shutil.rmtree(d, ignore_errors=True)
        d.mkdir(parents=True)
    pdf_jobs: list[tuple[Path, Path]] = []
    with tempfile.TemporaryDirectory() as tmpd:
        tmp = Path(tmpd)
        for meta in docs:
            body = json.loads((EXPANDED / "seed" / f"{meta['id']}.json").read_text())["body"]
            meta = {**meta, "ownerName": crew.get(meta.get("ownerId"), {}).get("name", "LineMate")}
            render(meta, body, SEED_OUT / meta["fileName"], pdf_jobs, tmp)
        manifest = []
        for f in sorted((EXPANDED / "samples").glob("SMP-*.json")):
            s = json.loads(f.read_text())
            name = f"{slug(s['title'])}.{s['fileKind']}"
            meta = {**s, "ownerName": crew.get(s.get("ownerId"), {}).get("name", "LineMate")}
            render(meta, s["body"], SAMPLE_OUT / name, pdf_jobs, tmp)
            manifest.append(
                {
                    "file": name,
                    "title": s["title"],
                    "category": s["category"],
                    "station": s["station"],
                    "ownerId": s.get("ownerId"),
                    "summary": s.get("summary", ""),
                    "tags": s.get("tags", []),
                }
            )
        # ticket attachment: chiller service log
        log_md = (
            "# Blast Chiller Service Log — Unit BC-2\n\n> Field service summary attached to TKT-017.\n\n"
            "## Readings\n\n| Time | Product | Core temp | Cabinet temp | Note |\n|---|---|---|---|---|\n"
            "| 14:05 | Chili, 3 in pans | 162°F | 31°F | Hard chill started |\n"
            "| 15:05 | Chili | 118°F | 38°F | Fan noise, E4 code shown |\n"
            "| 16:05 | Chili | 88°F | 44°F | Stage 1 limit missed (needs 70°F by 2 h) |\n\n"
            "## Technician findings\n\n- Error **E4**: evaporator fan motor drawing low current.\n"
            "- Condenser coil heavily soiled; airflow reduced.\n- Door gasket torn at lower hinge corner.\n\n"
            "## Actions\n\n1. Fan motor ordered (2 business days).\n2. Coil cleaned on site.\n"
            "3. Unit tagged *do not use for cooling* until the fan is replaced.\n"
        )
        render(
            {"id": "ATT-2", "title": "Blast Chiller Service Log", "fileKind": "pdf"},
            log_md,
            ATT_OUT / "chiller-service-log.pdf",
            pdf_jobs,
            tmp,
        )
        by_dir: dict[Path, list[Path]] = {}
        for src, _ in pdf_jobs:
            by_dir.setdefault(src.parent, []).append(src)
        for srcdir, items in by_dir.items():
            docx_to_pdf(items, srcdir)
        for src, dest in pdf_jobs:
            shutil.move(str(src.with_suffix(".pdf")), dest)
    for src, name in [
        ("blast-chiller-e4.jpg", "blast-chiller-e4.jpg"),
        ("fryer-oil.jpg", "fryer-oil-tuesday.jpg"),
        ("walk-in-shelf.jpg", "walk-in-shelf-6am.jpg"),
    ]:
        shutil.copy(SOURCE_PHOTOS / src, ATT_OUT / name)
    (SAMPLE_OUT / "samples.json").write_text(json.dumps(manifest, indent=2))
    for d in (SEED_OUT, SAMPLE_OUT, ATT_OUT):
        files = sorted(p for p in d.iterdir() if p.is_file())
        print(d.name, len(files), sum(p.stat().st_size for p in files) // 1024, "KB")


if __name__ == "__main__":
    main()
